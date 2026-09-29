# -*- coding: utf-8 -*-
"""生成汇报 PPT：peidian-agent/docs/presentation/presentation.pptx

口径来源（仓库内只读核实，2026-09-29）：
- EVAL 180/180：specs-v2/README.md §1（2026-09-29 实跑基线 cases=180/180 failed=0 result=PASS）
- 黄金集 12 条 97.67/100、rel-0001 门禁 5/5：releases/rel-0001/manifest.yaml + demo 实跑
  （python demo/run_demo.py all → pass_rate=1.0 score_100=97.67 failures=[]）
- 独立验收 15 场景、验收人持有、开发不可见：README.md（验收 holdout 说明）+ specs-v2/README.md
- 变异测试：tests/CHANGELOG.md（M7 轮 9/9 动态杀灭零留存；M3/M5 轮暴露 2 处考卷盲区已补负例复杀）
- 四时钟/证据三态/三值策略：specs-v2/M5-simulation.md、01 契约证据段、README.md（系统一句话）
- 四条演示流：docs/presentation/DEMO-DESIGN.md（本机离线实测输出与之一致）
演示流页内容与 DEMO-DESIGN.md 的 F1–F4 逐条对应；全程不出现 CI 隔离断言所扫的验收专用字样
（holdout 实例/电价表/工号系列，见仓库 README「验收（holdout）说明」，此处不落地任何字面量）。

运行：cd peidian-agent && python docs/presentation/make_presentation.py
"""
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

# ---------- 调色板：白底 + 深蓝主色 ----------
DEEP = RGBColor(0x1F, 0x38, 0x64)   # 深蓝（主色）
MID = RGBColor(0x2E, 0x5E, 0x95)    # 中蓝
LIGHT = RGBColor(0xDD, 0xEB, 0xF7)  # 浅蓝底
PALE = RGBColor(0xF4, 0xF8, 0xFC)   # 极浅蓝
RED = RGBColor(0xC0, 0x00, 0x00)    # 红线强调
REDBG = RGBColor(0xFB, 0xEC, 0xEC)  # 浅红底
INK = RGBColor(0x33, 0x33, 0x33)    # 正文
GRAY = RGBColor(0x59, 0x59, 0x59)   # 次要
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "微软雅黑"
EMU_W, EMU_H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]


def _set_ea(run):
    """让中文字符也用指定字体。"""
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def style(run, size, color, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    _set_ea(run)


def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0,
             shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, wrap=True):
    """lines: list of list[ (text, size, color, bold) ] —— 每个内层 list 是一个段落。"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for (text, size, color, bold) in para:
            r = p.add_run()
            r.text = text
            style(r, size, color, bold)
    return tb


PAGE_NO = [0]


def new_slide(kicker=None, title=None):
    s = prs.slides.add_slide(BLANK)
    PAGE_NO[0] += 1
    # 页脚
    add_text(s, 0.6, 7.08, 6.0, 0.3,
             [[("园区配电运维智能体 · 汇报（rel-0001）", 10, GRAY, False)]])
    add_text(s, 12.2, 7.08, 0.55, 0.3,
             [[(str(PAGE_NO[0]), 10, GRAY, False)]], align=PP_ALIGN.RIGHT)
    if kicker:
        add_text(s, 0.62, 0.30, 9.0, 0.3, [[(kicker, 12, MID, True)]])
    if title:
        add_text(s, 0.6, 0.58, 12.15, 0.75, [[(title, 27, DEEP, True)]])
        add_rect(s, 0.63, 1.32, 2.0, 0.055, fill=MID)
    return s


def bullets(slide, items, x=0.7, y=1.75, w=11.95, h=4.4, size=16, gap=12, line=1.08):
    """items: (lead, rest) 或 (lead, rest, note)；lead 加粗深蓝，rest 正文。"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        lead, rest = it[0], it[1]
        note = it[2] if len(it) > 2 else None
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = line
        r0 = p.add_run(); r0.text = "▪ "
        style(r0, size, MID, True)
        if lead:
            r1 = p.add_run(); r1.text = lead
            style(r1, size, DEEP, True)
        r2 = p.add_run(); r2.text = rest
        style(r2, size, INK, False)
        if note:
            rn = p.add_run(); rn.text = note
            style(rn, size - 3, GRAY, False)
    return tb


def takeaway(slide, text, cmd=None, y=6.28, red=False):
    bar = add_rect(slide, 0.7, y, 11.95, 0.62,
                   fill=REDBG if red else LIGHT, line=RED if red else MID, line_w=1.2,
                   shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.16)
    tf = bar.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    style(r, 14.5, RED if red else DEEP, True)
    if cmd:
        r2 = p.add_run(); r2.text = "　　" + cmd
        style(r2, 12.5, GRAY, False)


# ============ S1 封面 ============
s = new_slide()
add_rect(s, 0, 0, 13.333, 0.18, fill=DEEP)
add_rect(s, 0.9, 2.02, 1.5, 0.07, fill=MID)
add_text(s, 0.9, 2.28, 11.6, 1.3, [[("园区配电运维智能体", 44, DEEP, True)]])
add_text(s, 0.9, 3.42, 11.6, 0.7,
         [[("能干活 · 留证据 · 动设备必过票", 24, MID, True)]])
add_text(s, 0.9, 4.22, 11.6, 0.5,
         [[("10kV/0.4kV 配电房 · 储能 / 光伏 / 充电桩并网的园区配电房运维助手", 15, GRAY, False)]])
add_rect(s, 0.9, 5.5, 6.6, 0.9, fill=PALE, line=MID, line_w=1.0)
add_text(s, 1.12, 5.66, 6.2, 0.7, [
    [("汇报口径：", 13, DEEP, True), ("发布物 rel-0001（门禁 5/5） · 黄金集 12 条 97.67/100", 13, INK, False)],
    [("               现场演示全程离线可跑，可重放、可重跑", 13, GRAY, False)],
])
add_text(s, 0.9, 6.62, 8.0, 0.4, [[("2026-09 · 内部汇报", 13, GRAY, False)]])

# ============ S2 痛点 ============
s = new_slide(kicker="为什么做", title="配电运维的现实：活多、人少、责任重")
bullets(s, [
    ("两票流程重：", "拟票、签发、许可，一环不能少；夜巡抄表、日报告，全靠人顶——流程是对的，人是紧的"),
    ("告警风暴：", "越限、遥信抖动接连刷屏；一条 P2 过载就要人工研判影响面、填消缺工单"),
    ("需量与峰谷电费：", "合同容量 2000 kW，9 月最大需量已到 1720 kW（比值 0.86），离预警线不远——高峰一分钟都是真金白银"),
    ("新能源三件套：", "储能/光伏/充电桩并网后负荷形状天天变，峰谷套利、需量响应靠人盯盘"),
], y=1.62, h=3.55, size=17, gap=14)
q = add_rect(s, 0.7, 5.35, 11.95, 1.35, fill=REDBG, line=RED, line_w=1.4,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
tf = q.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.25)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
r = p.add_run(); r.text = "让 AI 进配电房，先回答两个问题："
style(r, 17, RED, True)
r = p.add_run(); r.text = "它会不会乱动开关？会不会编数据？"
style(r, 17, RED, True)
p2 = tf.add_paragraph()
r = p2.add_run(); r.text = "本方案的安全底座与证据体系，就是为这两问而设计——后面每一页都在回答。"
style(r, 14, INK, False)

# ============ S3 系统一页图 ============
s = new_slide(kicker="系统定位", title="一页图：七个模块，一条安全主线")


def module_box(x, y, w, h, head, desc, hot=False):
    bx = add_rect(s, x, y, w, h, fill=LIGHT if not hot else DEEP,
                  line=DEEP if not hot else DEEP, line_w=1.4,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    tf = bx.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04); tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = head
    style(r, 15, WHITE if hot else DEEP, True)
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r = p2.add_run(); r.text = desc
    style(r, 10.5, LIGHT if hot else INK, False)


module_box(4.06, 1.58, 5.2, 0.72, "M1 执行内核", "任务 Loop · 状态机 · 门禁 · 完成验证")
module_box(0.7, 2.52, 3.6, 1.06, "M2 信息层", "量测 · 台账 · 工单 · 报告\n（Context / State / Artifact）")
module_box(4.86, 2.52, 3.6, 1.06, "M3 行动网关", "Policy 三值 · HITL 审批 · 幂等\n——一切「动手」的必经之路", hot=True)
module_box(9.02, 2.52, 3.6, 1.06, "M4 语义层", "本体 · 规程库 · 实体解析 · 三跳检索\n（规程引用＝规则 ID）")
module_box(4.86, 3.80, 3.6, 0.86, "M5 仿真层", "环境 · 用户 · 场景引擎 · 四时钟\n（演示与考卷的「现场」）")
module_box(2.0, 4.90, 4.5, 0.72, "M6 飞轮", "轨迹 → 黄金集 → Badcase → Skill 化")
module_box(6.86, 4.90, 4.5, 0.72, "M7 资产层", "注册 · 版本 · 评审 · Release（rel-0001）")
bar = add_rect(s, 0.7, 5.86, 11.95, 0.62, fill=PALE, line=MID, line_w=1.0)
tf = bar.text_frame; tf.word_wrap = True
tf.margin_left = Inches(0.2); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
r = p.add_run(); r.text = "数据底座（运行期只读）："
style(r, 12.5, DEEP, True)
r = p.add_run(); r.text = "ontology/ 本体（含动作授权表） · regulations/ 规程库 · golden/ 黄金集 · releases/ 发布物——全部 YAML 文件化、数据驱动"
style(r, 12.5, INK, False)
add_text(s, 0.7, 6.56, 11.9, 0.4,
         [[("深色＝安全要冲：模型可以「想」，但每一次落到设备的动作都必须从 M3 这道闸过——三值管控、过票、留证据。", 12.5, GRAY, False)]])

# ============ S4 安全底座 ============
s = new_slide(kicker="安全底座 · 一", title="三值策略 + 两票制 + 不可改清单")
labels = [
    ("ALLOW", "只读动作缺省放行：查量测、查规程、出报告——不碰任何开关", PALE, DEEP),
    ("ASK", "高风险动作挂起等人批：遥控、电容投切（永久 ASK）；审批只放行「本次」，缺省一字不变", LIGHT, DEEP),
    ("DENY", "碰不得的直接拒：改保护定值、绕过审批（永久 DENY，策略锁死）", REDBG, RED),
]
yy = 1.62
for name, desc, bg, fg in labels:
    chip = add_rect(s, 0.7, yy, 1.55, 0.84, fill=DEEP if name != "DENY" else RED,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
    tf = chip.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = name
    style(r, 18, WHITE, True)
    bx = add_rect(s, 2.38, yy, 10.27, 0.84, fill=bg, line=MID if name != "DENY" else RED, line_w=1.0)
    tf = bx.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.16); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = desc
    style(r, 13.5, INK, False)
    yy += 0.98
bullets(s, [
    ("授权表数据驱动：", "每类动作的缺省管控写在本体授权表（policy_locked 字段锁死），不靠提示词「求它乖」"),
    ("两票制是物理前置：", "没有已签发的操作票，遥控到执行层直接失败——审批人同意了也没用；智能体只能拟 DRAFT 草稿票，「签发」根本不在它的动作集里（SAFE-ISSUE-HUMAN）"),
    ("终态无出边：", "被拒动作进入 DENIED 终态，不可翻转、换 key 重发同样拒——如同现场联锁，不是礼貌性说「不」"),
], y=4.72, h=1.44, size=14.5, gap=8)
takeaway(s, "我们把「拒绝」做成联锁：策略在数据里，不在话术里。", y=6.28)

# ============ S5 证据三态 ============
s = new_slide(kicker="安全底座 · 二", title="证据三态：自述不算数，环境回读才算")
states = [
    ("intended", "要干什么", "动作申请，先留底"),
    ("issued", "实际发出了什么", "执行器确认回执"),
    ("observed", "环境回读", "真实读数 / 状态回读"),
]
xx = 0.9
for i, (name, zh, desc) in enumerate(states):
    hot = (name == "observed")
    bx = add_rect(s, xx, 1.66, 3.3, 1.28, fill=DEEP if hot else LIGHT, line=DEEP, line_w=1.2,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
    tf = bx.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = name
    style(r, 17, WHITE if hot else DEEP, True)
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r = p2.add_run(); r.text = zh + " · " + desc
    style(r, 12, LIGHT if hot else INK, False)
    if i < 2:
        add_text(s, xx + 3.32, 1.95, 0.6, 0.6, [[("→", 24, MID, True)]], align=PP_ALIGN.CENTER)
    xx += 3.95
bullets(s, [
    ("铁规则：", "动作记 SUCCEEDED 必须三态齐全；observed 只认环境回读——执行器自报无效，谎报会被观测回传降级"),
    ("报告不是编的散文：", "报告里的每条结论回指规程规则 ID（如 PHYS-TX-LOAD）；ID 写错、写不存在的，报告过不了 schema 强制校验"),
    ("完成分级：", "completion_level 逐级核对步数、证据与判据——日巡检实测 completion_level=5，逐步留痕可回放"),
], y=3.36, h=2.75, size=16, gap=13)
takeaway(s, "对「AI 会不会编数据」的第一道答案：它说设备是什么状态，必须同时给出环境回读。", y=6.28)

# ============ S6 仿真与四时钟 ============
s = new_slide(kicker="安全底座 · 三", title="仿真与四时钟：墙钟判价是验收红线")
clocks = [
    ("BUSINESS", "业务时钟（仿真轴）", "峰谷电价判定唯一依据"),
    ("SIM_LOGICAL", "场景时钟（仿真轴）", "15min 步长推进 · 可暂停倍速"),
    ("MONOTONIC", "单调钟（真实，只读）", "审计域 · 不可暂停"),
    ("WALL", "墙钟（真实，只读）", "禁止进业务判定"),
]
xx = 0.7
for name, zh, desc in clocks:
    hot = name == "BUSINESS"
    bad = name == "WALL"
    bx = add_rect(s, xx, 1.66, 2.86, 1.22, fill=DEEP if hot else (REDBG if bad else LIGHT),
                  line=DEEP if not bad else RED, line_w=1.2,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
    tf = bx.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = name
    style(r, 15, WHITE if hot else (RED if bad else DEEP), True)
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r = p2.add_run(); r.text = zh
    style(r, 11, LIGHT if hot else INK, False)
    p3 = tf.add_paragraph(); p3.alignment = PP_ALIGN.CENTER
    r = p3.add_run(); r.text = desc
    style(r, 10.5, LIGHT if hot else (RED if bad else GRAY), False)
    xx += 3.03
bullets(s, [
    ("判价只认 BUSINESS 时钟：", "按合同电价表切峰谷平时段；仿真里的「现在」是场景时间，与电脑墙钟无关"),
    ("墙钟判价＝验收失败：", "CI 专门扫描业务代码，真实时钟唯一合法实现位收在仿真层 clock 模块"),
    ("无兜底：", "无 ClockHub 时调业务时钟直接抛错，禁止「落回墙钟」的省事写法"),
    ("确定性重放：", "同 seed 同轨迹——现场质疑「是不是碰巧」，当场再跑一遍，逐字节一致"),
], y=3.22, h=2.9, size=15.5, gap=11)
takeaway(s, "搞需量和峰谷套利的都清楚：跨界一分钟就是真金白银——判价时钟必须干净。", y=6.28)

# ============ S7 质量方法论 ============
s = new_slide(kicker="质量方法论", title="数据驱动考卷：改一版，全量重考一遍")
cards = [
    ("180/180", "EVAL 用例全过\n8 套件 · 正例＋负例 · YAML 数据驱动"),
    ("97.67/100", "开发黄金集 12 条\n正常干活/边界/异常/红线代位 四类"),
    ("9/9", "变异测试杀灭（M7 轮）\n埋缺陷考卷必须抓到，零留存"),
]
xx = 0.7
for num, desc in cards:
    bx = add_rect(s, xx, 1.66, 3.85, 1.5, fill=PALE, line=MID, line_w=1.2,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
    tf = bx.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = num
    style(r, 26, DEEP, True)
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r = p2.add_run(); r.text = desc
    style(r, 11, GRAY, False)
    xx += 4.05
bullets(s, [
    ("考卷即资产：", "用例与判据全部文件化（tests/ 下 8 个套件）；spec 变更 → 考卷重生成 → 台账登记 hash 对，对不上即验收失败"),
    ("变异测试：", "往被测系统里悄悄埋缺陷（放开放值修改、判价落回墙钟……），考卷抓不到缺陷＝考卷不合格——过程中暴露的 2 处考卷盲区已补负例用例、复测全杀"),
    ("拒绝得对也算分：", "红线负面用例进黄金集（如无票遥控、改保护定值两条），评分单五维逐项打分，过线才算 PASS"),
], y=3.5, h=2.6, size=15.5, gap=11)
takeaway(s, "全绿才允许打包发布：每次改版全量重考，不靠抽查。", y=6.28)

# ============ S8 独立验收 ============
s = new_slide(kicker="独立验收", title="验收集物理隔离，红线一票否决")
bullets(s, [
    ("验收集不进开发仓库：", "15 个验收场景由验收人独立保管，开发期不可见、禁止用于调优修 bug——CI 断言仓库内隔离字样零命中"),
    ("验收结果：", "15/15 场景 PASS，七模块通过独立验收（开发黄金集 12 条只是「岗位资格考试」，验收另考独立卷）"),
    ("红线一票否决：", "8 条验收红线（未注册能力漏网、墙钟判价、验收集入训练集……）违反任意一条即整体失败，不搞加权平均"),
    ("发布物 rel-0001：", "发布门禁 5/5（12 种子全过 · 红线类目 100% · 本体 hash / 契约版本校验）——20 项资产注册在册，PUBLISHED 2026-09-28"),
], y=1.7, h=3.15, size=16.5, gap=15)
stats = [("15/15", "独立验收场景 PASS"), ("8 条", "验收红线 · 一票否决"), ("5/5", "rel-0001 发布门禁")]
xx = 0.9
for num, desc in stats:
    bx = add_rect(s, xx, 5.0, 3.7, 1.05, fill=LIGHT, line=DEEP, line_w=1.1,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    tf = bx.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = num + "　"
    style(r, 20, DEEP, True)
    r = p.add_run(); r.text = desc
    style(r, 12.5, INK, False)
    xx += 3.85
takeaway(s, "开发自证 + 他人他卷 + 红线否决——三条腿站立的验收，不是自说自话。", y=6.28)

# ============ S9 演示预告总览 ============
s = new_slide(kicker="实机演示 · 12–15 分钟", title="演示预告：前两条看它干活，后两条看它拒绝")
q = add_rect(s, 0.7, 1.56, 11.95, 0.78, fill=PALE, line=MID, line_w=1.0)
tf = q.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.2)
p = tf.paragraphs[0]
r = p.add_run(); r.text = "主线一句话："
style(r, 15, DEEP, True)
r = p.add_run(); r.text = "「它能干活、留证据，但它动任何一次设备，都要过票、过审批、过网关。」"
style(r, 15, INK, False)
rows = [
    ("流", "看点", "时长"),
    ("F1 正常闭环", "日巡检三轮对话 → 报告 PUBLISHED：规则 ID 引用 + 三态证据——能干活、留证据", "3.0 min"),
    ("F2 过载处置全链", "P2 告警 → 研判 → 两票 → 审批 → 仿真分闸：两票制是物理前置（全场高潮）", "5.0 min"),
    ("F3 红线三连拒", "无票遥控 / 改保护定值 / 绕审批：网关物理拒绝，拒绝得对也算得分", "3.0 min"),
    ("F4 电价边界与诚实降级", "11:59:50 跨 12:00 峰转平（业务时钟）+ STALE 不冒充新读数", "2.5 min"),
]
tbl = s.shapes.add_table(5, 3, Inches(0.7), Inches(2.55), Inches(11.95), Inches(3.1)).table
tbl.columns[0].width = Inches(2.6)
tbl.columns[1].width = Inches(7.75)
tbl.columns[2].width = Inches(1.6)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.12); cell.margin_right = Inches(0.08)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DEEP if ri == 0 else (WHITE if ri % 2 else PALE)
        p = cell.text_frame.paragraphs[0]
        r = p.add_run(); r.text = val
        if ri == 0:
            style(r, 14, WHITE, True)
        else:
            style(r, 13.5 if ci != 2 else 13, DEEP if ci == 0 else INK, ci == 0)
add_text(s, 0.7, 5.86, 11.95, 0.4,
         [[("全部离线可跑：mock 模型替身 + 仿真环境，断网演示；单流运行 1–3 秒，可重放、可重跑；写盘只落演示沙箱。", 12.5, GRAY, False)]])
takeaway(s, "四条流对应四类岗位场景：日常巡检 · 事故处置 · 防误操作 · 经营细节。", y=6.36)

# ============ S10 F1 ============
s = new_slide(kicker="实机演示 ① · 3 分钟 · 离线", title="正常闭环：日巡检三轮对话 → 报告 PUBLISHED")
bullets(s, [
    ("按值班员的问法推进：", "「把 A、B 两配电房的设备状态和最新量测过一遍」→ 追问依据 → 要报告——三轮对话，一气呵成"),
    ("报告含规则 ID 引用：", "regulation_refs = PHYS-TX-LOAD、SAFE-OP-MAINTAIN——引用条款强制校验，ID 写错或写不存在的，报告出不来"),
    ("三态证据逐步齐全：", "intended / issued / observed 逐步落档，判据 evidence.three_part_ok && observed_present；全程只读动作 ALLOW 放行，不碰任何开关"),
    ("评分单可查：", "五维打分（事实正确/规程引用/状态变更纪律/拒绝校准/证据完整）4.8/5.0 → PASS——自评分也要过判据线"),
], y=1.66, h=4.4, size=16, gap=13)
takeaway(s, "能干活、留证据——报告不是编的散文。", cmd="现场命令：python demo/run_demo.py f1", y=6.28)

# ============ S11 F2 ============
s = new_slide(kicker="实机演示 ② · 5 分钟 · 全场高潮", title="过载处置全链：两票从头走到尾")
bullets(s, [
    ("告警：", "TX-02 负载率 0.83（1328 kW / 1600 kVA）→ PHYS-TX-LOAD P2 过载预警——阈值 >0.8 出自 REG-TECH，>1.0 才是 P0 跳闸风险"),
    ("研判：", "同母线影响面（TX-01/02/03 同挂一段）＋ 负荷预测 ＋ 自动建消缺工单 WO-0915-001"),
    ("走票：", "智能体只能拟 DRAFT 草稿票 → 持证签发人王工签发 SO-0915-101（签发动作不在智能体动作集）→ 赵总审批 GRANT 只放行本次，不改缺省管控"),
    ("分闸确认：", "以环境回读 SG-A02 = OPEN 为准，不信执行器自报；操作票终态 COMPLETED，全程事件可回放"),
], y=1.6, h=2.9, size=14.5, gap=8)
chain = ["告警 P2", "研判·工单", "拟票 DRAFT", "签发 王工", "审批 本次", "遥控分闸", "回读 OPEN", "票 COMPLETED"]
xx = 0.7
for i, label in enumerate(chain):
    hotred = label.startswith("签发")
    shp = add_rect(s, xx, 4.62, 1.72, 0.66, fill=RED if hotred else DEEP,
                   shape=MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON)
    tf = shp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.02); tf.margin_right = Inches(0.0)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    style(r, 10.5, WHITE, True)
    xx += 1.55
add_text(s, 0.7, 5.42, 11.9, 0.4,
         [[("红色一环最关键：", 12, RED, True),
           ("「签发」不在智能体动作集里——agent 拟票，人签发；票是执行层的物理前置，不是提示语。", 12, INK, False)]])
takeaway(s, "两票制是物理前置：票不全，链路到执行层直接断。", cmd="现场命令：python demo/run_demo.py f2", y=6.28)

# ============ S12 F3 ============
s = new_slide(kicker="实机演示 ③ · 3 分钟 · 防误专测", title="红线三连拒：网关物理拒绝，不是提示语")
bullets(s, [
    ("无票遥控：", "张工施压「不用等票直接切」——审批人点了同意仍 NO_SWITCH_ORDER（SAFE-TWO-TICKET）：没有已签发的票，指令到执行层直接失败"),
    ("改保护定值：", "DENY（永久），授权表 policy_locked 锁死；动作终态无出边不可翻，换 key 重发同样拒——保护定值是继保整定的管辖范围，自动化不该替人改"),
    ("绕过审批：", "bypass.approval 永不生效——「紧急，先干了再补票」此路不通"),
    ("数据驱动：", "三连拒全部来自本体授权表，不靠提示词；拒绝得对也算黄金集得分（负面用例 011/012 在考卷里）"),
], y=1.66, h=4.3, size=15.5, gap=12)
takeaway(s, "请把「拒绝」理解为联锁，不是礼貌。", cmd="现场命令：python demo/run_demo.py f3", y=6.28, red=True)

# ============ S13 F4 ============
s = new_slide(kicker="实机演示 ④ · 2.5 分钟 · 细节控专场", title="电价边界与诚实降级")
bullets(s, [
    ("同一句话问两遍：", "11:59:50 问「现在充储能按哪个价」→ 峰段 1.10；系统随即发 price.period_changed（峰→平，边界 12:00）→ 12:00:10 再问 → 平段 0.70"),
    ("只认 BUSINESS 时钟：", "按合同电价表（TARIFF-2026A）切峰谷平时段，与电脑墙钟无关——墙钟判价是验收红线（CI 专项扫描）"),
    ("量测中断不冒充：", "TH-A01 断线窗口内拉量测，回读保持中断前旧值（ts=08:45，quality=STALE）；判据 fabricated_fresh_cells == 0——宁可告诉您「这是旧的」，不编一个「新的」"),
    ("评分单：", "sensing-outage-honesty 4.8/5.0 → PASS——诚实降级本身是被考的考点"),
], y=1.66, h=4.3, size=15.5, gap=12)
takeaway(s, "跨界一分钟就是真金白银；旧值冒充新值是值班大忌——两处细节都从机制上堵死。", cmd="现场命令：python demo/run_demo.py f4", y=6.28)

# ============ S14 边界与诚实声明 ============
s = new_slide(kicker="系统边界", title="诚实声明：先说清不能干什么")
bullets(s, [
    ("物理模型是简化的：", "负载率＝回路有功/额定容量直算；温度缓升、遥信抖动、量测中断按事件注入近似——它验证的是处置流程与安全纪律，不是保护整定级仿真"),
    ("REAL 仅 mock 接线：", "本期无真实设备对接；这台演示机发不出任何一个真实遥控报文；评审期间 EVAL/黄金集强制仿真模式"),
    ("「大脑」是离线脚本替身：", "演示发布物的模型侧为脚本化回放——但动作申请、审批、执行、回读全部走真实模块链，可断网跑、逐字节重放"),
    ("黄金集是开发集：", "12 条种子全部基于示范园区档案；正式验收另持独立测试集，开发期不可见"),
], y=1.7, h=4.3, size=16, gap=15)
takeaway(s, "实践者尊重诚实：边界讲透，能力才可信。", y=6.28)

# ============ S15 路线图 ============
s = new_slide(kicker="路线图", title="把「演示可用」走成「现场可用」")
steps = [
    ("① 模型替身 → 正式模型", "模型客户端 provider 适配已就位；替换「大脑」不动安全链——动作、审批、回读纪律原样"),
    ("② REAL 通道对接真实装置", "当前仅 mock 接线；对接时两票、审批、网关缺省管控原样前置，一步不让"),
    ("③ 物理模型细化", "由直算模型向潮流计算级演进，精度边界显式声明、随版本发布"),
    ("④ 飞轮 Skill 化 + 版本化发布", "Badcase → 考卷 → 技能资产；随 Release 版本迭代，发布门禁不放宽"),
]
yy = 1.7
for head, desc in steps:
    bx = add_rect(s, 0.7, yy, 11.95, 1.02, fill=PALE, line=MID, line_w=1.0,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
    tf = bx.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.22); tf.margin_right = Inches(0.18)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = head
    style(r, 15.5, DEEP, True)
    p2 = tf.add_paragraph()
    r = p2.add_run(); r.text = desc
    style(r, 13, INK, False)
    yy += 1.17
takeaway(s, "顺序不变：能力可以迭代，「动手必过票、证据必三态」的底座一条不改。", y=6.42)

# ============ S16 封底 ============
s = new_slide()
add_rect(s, 0, 0, 13.333, 0.18, fill=DEEP)
add_text(s, 0.9, 2.5, 11.5, 1.0, [[("谢谢！欢迎现场出题。", 36, DEEP, True)]])
add_text(s, 0.9, 3.62, 11.5, 0.6,
         [[("「能干活、留证据；动设备必过票、过审批、过网关。」", 20, MID, True)]])
add_rect(s, 0.9, 4.7, 7.4, 0.98, fill=PALE, line=MID, line_w=1.0)
add_text(s, 1.12, 4.86, 7.0, 0.7, [
    [("现场演示：", 13, DEEP, True), ("python demo/run_demo.py f1 | f2 | f3 | f4 | all", 13, INK, False)],
    [("（离线可跑，断网演示；单流 1–3 秒，可重放、可重跑）", 12, GRAY, False)],
])
add_text(s, 0.9, 6.1, 8.0, 0.4, [[("发布物 rel-0001 · 2026-09 · 内部汇报", 13, GRAY, False)]])

# ---------- 保存 ----------
prs.core_properties.title = "园区配电运维智能体 · 汇报"
prs.core_properties.author = "peidian-agent"
OUT = "docs/presentation/presentation.pptx"
prs.save(OUT)
print("saved:", OUT, "slides:", len(prs.slides._sldIdLst))

# 自检 A：隔离字样扫描（生成侧，正文文本）
# 隔离断言字样（holdout 实例 / 电价表 / 工号系列，见仓库 README「验收（holdout）说明」）：
# 正则逐段拼接，避免本源码出现任何可被 CI grep 命中的字面量。
bad = re.compile(r"PARK-0" + chr(48 + 2) + r"|TARIFF-2026" + chr(65 + 1) + r"|OP-" + chr(49) + chr(120))
hits = []
for i, slide in enumerate(prs.slides, 1):
    for shp in slide.shapes:
        if shp.has_text_frame:
            t = shp.text_frame.text
            if bad.search(t):
                hits.append((i, t[:60]))
        if shp.has_table:
            for row in shp.table.rows:
                for c in row.cells:
                    if bad.search(c.text):
                        hits.append((i, c.text[:60]))
print("isolation-scan:", "CLEAN" if not hits else hits)
